"""Excel report styling configurations and layout templates."""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# Curated harmonious color palette and fonts
FONT_FAMILY = "Segoe UI"

# Fills
HEADER_FILL = PatternFill(
    start_color="FF1F4E79", end_color="FF1F4E79", fill_type="solid"
)  # Sleek dark blue
ZEBRA_FILL = PatternFill(
    start_color="FFF2F6FA", end_color="FFF2F6FA", fill_type="solid"
)  # Extremely light blue-gray
SUMMARY_LABEL_FILL = PatternFill(
    start_color="FFE9EEF4", end_color="FFE9EEF4", fill_type="solid"
)
ALERT_CRITICAL_FILL = PatternFill(
    start_color="FFFCE4D6", end_color="FFFCE4D6", fill_type="solid"
)  # Light orange/red for criticals

# Fonts
HEADER_FONT = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_FAMILY, size=16, bold=True, color="1F4E79")
SECTION_FONT = Font(name=FONT_FAMILY, size=13, bold=True, color="2C3E50")
STANDARD_FONT = Font(name=FONT_FAMILY, size=10, bold=False, color="000000")
BOLD_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
MUTED_FONT = Font(name=FONT_FAMILY, size=9, italic=True, color="7F8C8D")

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Borders
BORDER_THIN = Border(
    left=Side(style="thin", color="BDC3C7"),
    right=Side(style="thin", color="BDC3C7"),
    top=Side(style="thin", color="BDC3C7"),
    bottom=Side(style="thin", color="BDC3C7"),
)
BORDER_TOP_THICK = Border(top=Side(style="medium", color="1F4E79"))
BORDER_BOTTOM_DOUBLE = Border(bottom=Side(style="double", color="1F4E79"))
